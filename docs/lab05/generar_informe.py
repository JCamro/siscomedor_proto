"""Genera docs/lab05/INFORME_LAB05_IHC.docx a partir de la plantilla "LAB05 IHC.docx".

Uso (desde la raiz del repositorio):  python docs/lab05/generar_informe.py
Requiere: python-docx, matplotlib. Reejecutar tras agregar capturas en docs/lab05/capturas/.
El informe cubre solo la tarea principal "Comprar ticket".
"""
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "docs" / "lab05"
CAPS = OUT_DIR / "capturas"
TEMPLATE = ROOT / "LAB05 IHC.docx"
OUTPUT = OUT_DIR / "INFORME_LAB05_IHC.docx"
DIAGRAM = OUT_DIR / "diagrama_navegacion.png"

CAPTURAS = [
    ("01_login_vacio", "Login recién abierto con el botón «Ingresar» deshabilitado."),
    ("02_login_error_campos", "Mensajes de validación bajo el código y la contraseña incompletos."),
    ("03_login_credenciales_incorrectas", "Aviso de error ante credenciales incorrectas."),
    ("04_login_valido", "Credenciales válidas, contraseña visible y switch «Recordar mi código» activo."),
    ("05_inicio", "Pantalla de Inicio tras el ingreso, sin el login en la pila."),
    ("06_inicio_seleccion", "Variante de menú seleccionada y barra inferior con el total."),
    ("07_tu_seleccion", "Panel «Tu selección» con el botón «Revisar y pagar»."),
    ("08_checkout_confirmacion", "Diálogo «Revisar y pagar» antes de confirmar."),
    ("09_checkout_exito", "Resultado exitoso del pago."),
    ("10_mis_tickets", "Lista «Mis tickets» con el ticket comprado."),
    ("11_detalle_ticket", "Detalle del ticket en la ruta tickets/{ticketId}."),
    ("12_sin_conexion", "Inicio sin conexión con el pago bloqueado."),
    ("13_compra_duplicada", "Prevención de compra duplicada con «Ya tienes este ticket»."),
    ("14_tests_ok", "Pruebas unitarias en verde en Android Studio, incluida LoginValidationTest."),
    ("15_preview_login", "Panel @Preview de LoginPreviews.kt."),
]

# Indicación de qué capturar, visible en el documento mientras falte la imagen.
COMO_CAPTURAR = {
    "01_login_vacio": "Abrir la app recién instalada; no escribir nada.",
    "02_login_error_campos": "Escribir código 123 y contraseña ab; se ven los errores en rojo.",
    "03_login_credenciales_incorrectas": "Código 20240001, contraseña incorrecta1, pulsar Ingresar.",
    "04_login_valido": "Código 20240001, contraseña comedor123, tocar el ojo y activar el switch.",
    "05_inicio": "Pulsar Ingresar con credenciales correctas.",
    "06_inicio_seleccion": "En Inicio, elegir fecha, servicio y tocar una variante de menú.",
    "07_tu_seleccion": "Tocar el resumen de la barra verde inferior.",
    "08_checkout_confirmacion": "Pulsar «Revisar y pagar» o «Pagar» en la barra verde.",
    "09_checkout_exito": "Pulsar «Confirmar pago» en el diálogo.",
    "10_mis_tickets": "Cerrar el diálogo y tocar la pestaña «Mis tickets».",
    "11_detalle_ticket": "Tocar un ticket de la lista.",
    "12_sin_conexion": "En Inicio, cambiar la conexión demo a Sin conexión.",
    "13_compra_duplicada": "Elegir la cena de hoy, que ya está comprada.",
    "14_tests_ok": "Ejecutar los tests unitarios en Android Studio y capturar el panel Run.",
    "15_preview_login": "Abrir LoginPreviews.kt en vista Split.",
}


# ---------------------------------------------------------------- diagrama
def draw_diagram(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(13, 7.6), dpi=160)
    ax.set_xlim(0, 13); ax.set_ylim(0, 7.6); ax.axis("off")
    W, H = 2.3, 1.0
    nodes = {
        "login": (1.5, 5.0, "1. Login\nruta: login", "screen"),
        "home": (4.6, 5.0, "2. Inicio\nfecha, servicio y variante", "screen"),
        "sel": (7.8, 5.0, "3. Tu selección\npanel sobre Inicio", "overlay"),
        "rev": (11.2, 5.0, "4. Revisar y pagar\ndiálogo de confirmación", "dialog"),
        "res": (11.2, 2.0, "5. Resultado\nmismo diálogo", "dialog"),
        "tick": (7.8, 2.0, "6. Mis tickets\nruta: tickets", "screen"),
        "det": (4.6, 2.0, "7. Detalle del ticket\nruta: tickets/{ticketId}", "screen"),
    }
    colors = {"screen": ("#DDEBF7", "-"), "overlay": ("#FFF2CC", "-"), "dialog": ("#FCE4D6", "--")}
    for x, y, label, kind in nodes.values():
        fc, ls = colors[kind]
        ax.add_patch(FancyBboxPatch((x - W / 2, y - H / 2), W, H, boxstyle="round,pad=0.06",
                                    fc=fc, ec="#1F4E79", lw=1.6, ls=ls))
        ax.text(x, y, label, ha="center", va="center", fontsize=10)

    def arrow(p1, p2, text, tx, ty, back=False, rad=0.0):
        color = "#C00000" if back else "#1F4E79"
        ax.annotate("", xy=p2, xytext=p1,
                    arrowprops=dict(arrowstyle="-|>", color=color, lw=1.6, mutation_scale=16,
                                    connectionstyle=f"arc3,rad={rad}", ls="--" if back else "-"))
        ax.text(tx, ty, text, fontsize=9, ha="center", va="center", color=color,
                bbox=dict(fc="white", ec="none", pad=1.5))

    # Avance (azul, continuo)
    arrow((2.7, 5.15), (3.4, 5.15), "Ingresar\n(credenciales válidas)", 3.05, 5.95)
    arrow((5.8, 5.15), (6.6, 5.15), "Tocar el\nresumen", 6.2, 5.95)
    arrow((9.0, 5.15), (10.0, 5.15), "Pulsar\nRevisar y pagar", 9.5, 5.95)
    arrow((4.6, 5.55), (11.2, 5.55), "Atajo: pulsar Pagar en la barra verde", 7.9, 6.75, rad=-0.28)
    arrow((11.2, 4.45), (11.2, 2.55), "Pulsar\nConfirmar pago", 12.35, 3.5)
    arrow((10.0, 2.15), (9.0, 2.15), "Cerrar y tocar\npestaña Mis tickets", 9.5, 2.95)
    arrow((6.6, 2.15), (5.8, 2.15), "Tocar un\nticket", 6.2, 2.95)
    # Regreso (rojo, discontinuo)
    arrow((6.6, 4.75), (5.8, 4.75), "Atrás", 6.2, 4.5, back=True)
    arrow((10.0, 4.6), (5.8, 4.6), "Atrás / cerrar diálogo", 9.3, 3.75, back=True, rad=-0.3)
    arrow((5.8, 1.85), (6.6, 1.85), "Atrás", 6.2, 1.4, back=True)
    arrow((7.8, 2.55), (5.2, 4.45), "Atrás", 6.75, 3.3, back=True)

    ax.text(0.3, 0.75, "Nota: al ingresar, Login se elimina de la pila; Atrás desde Inicio cierra la app.",
            fontsize=10, ha="left", weight="bold", color="#1F4E79")
    ax.text(0.3, 0.3, "Azul continuo: avanzar.   Rojo discontinuo: Atrás.   Celeste: pantalla con ruta.   "
            "Amarillo: panel.   Naranja con borde discontinuo: diálogo de confirmación.",
            fontsize=9, ha="left", style="italic")
    ax.set_title("SisCome: flujo de la tarea «Comprar ticket»", fontsize=14, weight="bold")
    fig.tight_layout(); fig.savefig(path); plt.close(fig)


# ---------------------------------------------------------------- helpers
def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def para(cell, text="", bold=False, size=10, italic=False, align=None, mono=False):
    p = cell.add_paragraph()
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if mono:
            r.font.name = "Consolas"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(4)
    return p


def text(cell, t):
    return para(cell, t, align=WD_ALIGN_PARAGRAPH.JUSTIFY)


def heading(cell, t):
    p = para(cell, t, bold=True, size=11)
    p.runs[0].font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(8)


def caption(cell, t):
    para(cell, t, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)


def code(cell, t):
    for line in t.strip("\n").split("\n"):
        p = para(cell, line, size=8.5, mono=True)
        p.paragraph_format.space_after = Pt(0)
    para(cell, "")


def table(cell, header, rows, widths):
    t = cell.add_table(rows=1, cols=len(header))
    t.style = "Table Grid" if "Table Grid" in [s.name for s in cell.part.document.styles] else None
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9)
        shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(v); r.font.size = Pt(8.5)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    if t.style is None:  # bordes manuales si la plantilla no trae "Table Grid"
        tblPr = t._tbl.tblPr; borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "808080")
            borders.append(e)
        tblPr.append(borders)
    para(cell, "")
    return t


def clear_trailing(cell):
    for p in list(cell.paragraphs)[1:]:
        if not p.text.strip():
            p._element.getparent().remove(p._element)


def find_cell(doc, title):
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                if c.text.strip().startswith(title):
                    return c
    raise SystemExit(f"No se encontró la sección: {title}")


# ---------------------------------------------------------------- contenido
def ejercicios(cell, missing):
    clear_trailing(cell)
    text(cell, "Esta sección documenta las siete fases de la guía aplicadas a una sola tarea de SisCome: comprar un "
               "ticket del comedor. Solo se describen las pantallas, eventos y mensajes que intervienen en esa "
               "tarea, desde el ingreso hasta la consulta del ticket comprado.")

    heading(cell, "Fase 1. Descripción de la aplicación y tarea principal")
    text(cell, "SisCome es una aplicación Android desarrollada con Kotlin y Jetpack Compose, la biblioteca de Android "
               "que construye la interfaz a partir de funciones. Permite a los estudiantes universitarios adquirir "
               "tickets para el comedor desde el celular, normalmente en momentos breves entre clases y con "
               "conectividad variable.")
    text(cell, "La tarea principal es comprar un ticket y sigue una secuencia fija de siete pasos:")
    for i, paso in enumerate([
        "Ingresar con el código universitario y la contraseña en Login.",
        "Elegir en Inicio la fecha, el servicio y la variante de menú.",
        "Revisar lo elegido y el total en el panel «Tu selección».",
        "Comprobar el resumen en el diálogo «Revisar y pagar» y confirmar.",
        "Leer el resultado del pago en el mismo diálogo.",
        "Abrir la lista «Mis tickets» desde la barra inferior.",
        "Tocar el ticket para ver su detalle con el código QR.",
    ], 1):
        para(cell, f"{i}. {paso}")

    heading(cell, "Fase 2. Diagrama del flujo de navegación")
    text(cell, "La navegación se implementó con Navigation Compose. Una ruta es el nombre de texto de una pantalla, "
               "como una dirección web interna; en la compra intervienen login, home, tickets y tickets/{ticketId}, "
               "donde {ticketId} se reemplaza por el identificador del ticket elegido. Las rutas se declaran en la "
               "sealed class Route de ui/navigation/Routes.kt, una lista cerrada de nombres de pantalla que impide "
               "navegar a una pantalla que no existe.")
    text(cell, "El NavHost es el contenedor que muestra la pantalla correspondiente a la ruta activa, y el "
               "NavController es el objeto que cambia de pantalla y recuerda el recorrido. Ese recorrido se guarda en "
               "la pila de navegación o back stack: cada pantalla nueva se apila encima de la anterior y el botón "
               "Atrás retira la de arriba. El panel «Tu selección» y el diálogo de pago no son rutas, sino ventanas "
               "que se dibujan encima de Inicio; por eso Atrás simplemente las cierra.")
    text(cell, "Como se ilustra en la Figura 1, las flechas azules numeradas recorren la tarea completa y las flechas "
               "rojas indican el regreso con Atrás. El diagrama muestra que el flujo no tiene callejones sin salida y "
               "que el login desaparece de la pila al ingresar, de modo que Atrás desde Inicio cierra la aplicación "
               "en lugar de volver al formulario.")
    cell.add_paragraph().add_run().add_picture(str(DIAGRAM), width=Cm(16))
    caption(cell, "Figura 1. Flujo de navegación de la tarea «Comprar ticket».")

    heading(cell, "Fase 3. Gestión de eventos")
    text(cell, "Un evento es una acción del usuario, como tocar o escribir, a la que la aplicación responde. En "
               "Compose cada componente recibe una función que se ejecuta ante su evento: onClick al tocarlo, "
               "onValueChange cada vez que cambia el texto de un campo y onCheckedChange al activar o desactivar un "
               "switch. La Tabla 1 resume los eventos de la compra; se observa que todos actualizan el estado de la "
               "pantalla o cambian de ruta, y ninguno cobra sin pasar por la confirmación.")
    caption(cell, "Tabla 1. Eventos implementados en la tarea «Comprar ticket».")
    table(cell, ["Componente", "Evento", "Archivo", "Comportamiento"], [
        ["Campo «Código universitario»", "onValueChange", "ui/login/LoginScreen.kt", "Conserva solo dígitos hasta 8, guarda el valor y lo valida con validateCode()."],
        ["Campo «Contraseña»", "onValueChange", "ui/login/LoginScreen.kt", "Guarda el valor y exige al menos 6 caracteres con validatePassword()."],
        ["Icono mostrar u ocultar", "onClick", "ui/login/LoginScreen.kt", "Alterna la visibilidad de la contraseña."],
        ["Switch «Recordar mi código»", "onCheckedChange", "ui/login/LoginScreen.kt", "Activa o desactiva la preferencia."],
        ["Botón «Ingresar»", "onClick", "ui/login/LoginScreen.kt", "Verifica las credenciales y navega a Inicio o muestra un aviso."],
        ["Fecha, servicio y variante", "onClick", "ui/home/HomeScreen.kt", "Marca la opción elegida y actualiza el total de la barra verde."],
        ["Resumen de la barra verde", "onClick", "ui/selection/SelectionUi.kt", "Abre el panel «Tu selección»."],
        ["Botones «Revisar y pagar» y «Pagar»", "onClick", "ui/selection/SelectionUi.kt", "Abren el diálogo de revisión sin cobrar todavía."],
        ["Botón «Confirmar pago»", "onClick", "ui/checkout/CheckoutDialog.kt", "Procesa el pago y muestra el resultado."],
        ["Pestaña «Mis tickets»", "onClick", "ui/SisComeApp.kt", "Navega a la ruta tickets."],
        ["Ticket de la lista", "onClick", "ui/tickets/TicketsUi.kt", "Navega a tickets/{ticketId}."],
        ["Botón «Volver a Mis tickets»", "onClick", "ui/tickets/TicketsUi.kt", "Llama a popBackStack() y regresa a la lista."],
    ], [3.6, 2.4, 3.8, 6.2])

    heading(cell, "Fase 4. Feedback")
    text(cell, "El feedback es la respuesta visible con la que la aplicación confirma que entendió la acción. En el "
               "login, cada campo usa isError para pintarse de rojo cuando el dato no es válido y supportingText para "
               "mostrar debajo el mensaje de error o una ayuda. La Tabla 2 agrupa las respuestas por tipo y evidencia "
               "que cada mensaje indica la causa y el siguiente paso, no solo que algo falló.")
    caption(cell, "Tabla 2. Feedback ofrecido durante la compra.")
    table(cell, ["Tipo", "Situación", "Mensaje o respuesta visual"], [
        ["Ayuda", "Escritura del código", "Contador «n/8 dígitos» y ayuda «Mínimo 6 caracteres» para la contraseña."],
        ["Error", "Campos inválidos", "Texto en rojo bajo el campo, por ejemplo «El código debe tener 8 dígitos (llevas 3)»."],
        ["Error", "Credenciales incorrectas", "Aviso «No se pudo ingresar» con la causa."],
        ["Éxito", "Selección de menú", "Opción resaltada y barra verde con el número de servicios y el total."],
        ["Error", "Sin conexión", "Aviso de que no se puede iniciar un pago nuevo y botón de confirmación deshabilitado."],
        ["Pendiente", "Pago sin emisión confirmada", "Indicación de no pagar otra vez y de consultar la misma referencia."],
        ["Éxito", "Pago confirmado", "Resultado en el diálogo y ticket disponible en «Mis tickets»."],
    ], [2.4, 4.2, 9.4])

    heading(cell, "Fase 5. Prevención de errores")
    text(cell, "La prevención busca que el error no llegue a ocurrir, en vez de informarlo después. La Tabla 3 "
               "relaciona cada mecanismo con el error que evita; destaca que el botón «Ingresar» y el de "
               "confirmación permanecen deshabilitados mientras falte información o conexión, lo que elimina el "
               "intento fallido antes de que suceda.")
    caption(cell, "Tabla 3. Mecanismos de prevención de errores.")
    table(cell, ["Mecanismo", "Implementación", "Error que evita"], [
        ["Validación de campos", "validateCode() y validatePassword() en domain/LoginValidation.kt; el error aparece después de tocar el campo.", "Enviar datos incompletos o con formato inválido."],
        ["Filtro de entrada", "sanitizeCode() descarta lo que no es número y limita a 8 dígitos.", "Escribir letras o códigos demasiado largos."],
        ["Botón deshabilitado", "El botón «Ingresar» se habilita solo con datos válidos; «Confirmar pago» solo con conexión.", "Operar con información incompleta o sin red."],
        ["Compra duplicada", "Un servicio ya comprado muestra «Ya tienes este ticket» y ofrece el próximo menú disponible.", "Pagar dos veces el mismo ticket."],
        ["Confirmación previa", "El diálogo «Revisar y pagar» exige pulsar «Confirmar pago».", "Pagar por un toque accidental."],
        ["Pila de navegación", "popUpTo(login) con inclusive = true al ingresar.", "Volver al login con Atrás tras autenticarse."],
    ], [3.2, 7.4, 5.4])

    heading(cell, "Fase 6. Patrones de interacción")
    text(cell, "Un patrón de interacción es una solución conocida para un problema de uso frecuente. La compra combina "
               "tres patrones. El formulario con confirmación impide que el pago se ejecute con un solo toque: el "
               "diálogo «Revisar y pagar» presenta la selección y el total, espera la confirmación y luego comunica "
               "el resultado. El patrón lista-detalle organiza «Mis tickets», donde cada elemento abre su detalle "
               "mediante la ruta tickets/{ticketId}. Finalmente, la navegación de regreso garantiza que Atrás siempre "
               "lleve a un lugar predecible, como se observa en las flechas rojas de la Figura 1.")

    heading(cell, "Fase 7a. Fragmentos de código relevantes")
    text(cell, "Se presentan tres fragmentos breves de ui/SisComeApp.kt y domain/LoginValidation.kt, elegidos porque "
               "sostienen las decisiones de navegación y validación descritas en las fases anteriores.")
    text(cell, "El primer fragmento configura el NavHost. Cada llamada composable(ruta) asocia una ruta con la función "
               "que dibuja su pantalla; en Compose, un composable es precisamente una función que dibuja una parte de "
               "la interfaz. La función navigate(ruta) abre una pantalla nueva encima de la pila:")
    code(cell, """
NavHost(navController, startDestination = Route.Login.route) {
    composable(Route.Login.route) {
        LoginScreen(onLoginSuccess = {
            navController.navigate(Route.Home.route) {
                popUpTo(Route.Login.route) { inclusive = true }
            }
        })
    }
    composable(Route.TicketDetail.route, arguments = listOf(
        navArgument(Route.TicketDetail.ARG) { type = NavType.StringType })) { entry ->
        val ticketId = entry.arguments?.getString(Route.TicketDetail.ARG).orEmpty()
        TicketDetailScreen(ticketId, model, onBack = { navController.popBackStack() })
    }
}""")
    text(cell, "Al ingresar, popUpTo(login) retira de la pila las pantallas hasta el login, y inclusive = true indica "
               "que también se retire el propio login; por eso Atrás desde Inicio no regresa al formulario. En el "
               "detalle, el botón «Volver a Mis tickets» llama a popBackStack(), que quita la pantalla actual de la "
               "pila y muestra la anterior; el botón Atrás del sistema produce el mismo efecto.")
    text(cell, "El segundo fragmento es el que ejecuta la pestaña «Mis tickets» de la barra inferior:")
    code(cell, """
navController.navigate(target.route) {
    popUpTo(Route.Home.route) { saveState = true }
    launchSingleTop = true
    restoreState = true
}""")
    text(cell, "La pila se recorta hasta Inicio, de modo que Atrás desde «Mis tickets» vuelve a Inicio. "
               "launchSingleTop evita apilar dos veces la misma pantalla si se toca la pestaña repetidamente, y "
               "saveState junto con restoreState conservan lo que el usuario tenía en pantalla al regresar.")
    text(cell, "El tercer fragmento valida el código universitario. Es una función pura, es decir, solo recibe el "
               "texto y devuelve un mensaje de error o null si el dato es correcto, lo que permite probarla sin "
               "abrir la aplicación:")
    code(cell, """
fun validateCode(code: String): String? = when {
    code.isBlank() -> "Ingresa tu código universitario."
    !code.all { it.isDigit() } -> "El código solo admite números."
    code.length != CODE_LENGTH -> "El código debe tener $CODE_LENGTH dígitos (llevas ${code.length})."
    else -> null
}""")
    text(cell, "El mensaje devuelto se muestra con supportingText y activa isError en el campo, que es el feedback "
               "descrito en la Tabla 2.")

    heading(cell, "Fase 7b. Capturas de las interfaces")
    text(cell, "Las Figuras 2 a 16 recorren la compra en el mismo orden del diagrama de la Figura 1. Las Figuras 2 a 5 "
               "muestran la validación del login, las Figuras 6 a 12 el avance desde Inicio hasta el detalle del "
               "ticket, y las Figuras 13 y 14 los mecanismos de prevención sin conexión y de compra duplicada. La "
               "Figura 15 evidencia las pruebas unitarias en verde, y la Figura 16 corresponde a @Preview, la "
               "anotación que dibuja una pantalla dentro de Android Studio sin ejecutar la aplicación.")
    grid = cell.add_table(rows=0, cols=2)
    for i in range(0, len(CAPTURAS), 2):
        cells = grid.add_row().cells
        for j, (name, cap) in enumerate(CAPTURAS[i:i + 2]):
            c = cells[j]; img = CAPS / f"{name}.png"
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if img.exists():
                p.add_run().add_picture(str(img), width=Cm(7.5))
            else:
                missing.append(name)
                r = p.add_run(f"[Insertar captura: {name}]"); r.bold = True
                r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
                hp = c.add_paragraph(); hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = hp.add_run(f"Qué capturar: {cap} Cómo: {COMO_CAPTURAR[name]}")
                r.font.size = Pt(8); r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            cp = c.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cp.add_run(f"Figura {i + j + 2}. {cap}"); r.italic = True; r.font.size = Pt(8.5)
    para(cell, "")

    heading(cell, "Fase 7c. Matriz de pruebas")
    text(cell, "La Tabla 4 registra los casos ejecutados sobre la tarea de compra, incluidos el avance, el regreso y "
               "las situaciones de error. Todos los casos resultaron conformes, lo que indica que la navegación y la "
               "prevención de errores se comportan como se diseñaron.")
    caption(cell, "Tabla 4. Matriz de pruebas de la tarea «Comprar ticket».")
    table(cell, ["ID", "Caso", "Pasos", "Resultado esperado", "Resultado"], [
        ["P01", "Campos vacíos", "Abrir la app sin escribir datos.", "Botón «Ingresar» deshabilitado.", "Conforme"],
        ["P02", "Formato inválido", "Código de 3 dígitos y contraseña de 2 caracteres.", "Mensajes de error bajo cada campo.", "Conforme"],
        ["P03", "Caracteres no numéricos", "Escribir letras en el código.", "Las letras se descartan.", "Conforme"],
        ["P04", "Credenciales incorrectas", "Código 20240001 y contraseña errónea; pulsar Ingresar.", "Aviso «No se pudo ingresar».", "Conforme"],
        ["P05", "Ingreso correcto", "Credenciales de demostración; pulsar Ingresar.", "Se abre Inicio; Atrás no vuelve al login.", "Conforme"],
        ["P06", "Compra con confirmación", "Elegir variante, abrir «Tu selección», pulsar «Revisar y pagar» y «Confirmar pago».", "Diálogo de revisión y resultado exitoso.", "Conforme"],
        ["P07", "Consulta del ticket", "Tocar la pestaña «Mis tickets» y luego el ticket.", "Se abre tickets/{ticketId} con el ticket correcto.", "Conforme"],
        ["P08", "Navegación de regreso", "En el detalle pulsar Atrás; en «Mis tickets» pulsar Atrás.", "Vuelve a «Mis tickets» y luego a Inicio.", "Conforme"],
        ["P09", "Compra duplicada", "Elegir un servicio ya comprado.", "Se muestra «Ya tienes este ticket»; no se ofrece pagar.", "Conforme"],
        ["P10", "Sin conexión", "Simular estado sin red e intentar pagar.", "Aviso de red y «Confirmar pago» deshabilitado.", "Conforme"],
        ["P11", "Pruebas unitarias", "Ejecutar LoginValidationTest con sus 6 casos y el resto de la suite.", "Todas las pruebas en verde.", "Conforme"],
    ], [1.1, 3.0, 4.6, 5.3, 2.0])


def cuestionario(cell):
    clear_trailing(cell)
    text(cell, "Las respuestas se apoyan en la implementación de la tarea «Comprar ticket» descrita en la solución "
               "de ejercicios.")
    qa = [
        ("1. ¿Cuál es la diferencia entre diseñar una pantalla y diseñar un flujo de interacción completo?",
         ["Diseñar una pantalla se centra en la organización visual de un único estado de la interfaz: jerarquía, "
          "controles, textos y distribución. Diseñar un flujo define la secuencia de estados que el usuario recorre "
          "para alcanzar un objetivo, con las acciones que provocan cada transición, las respuestas del sistema, las "
          "cancelaciones y los errores (Díaz, Harari y Amadeo, 2013; ISO 9241-210, 2019).",
          "En SisCome, una pantalla de Inicio bien diseñada no garantiza la compra. Fue necesario definir qué ocurre "
          "después del login, cómo se confirma el pago, qué sucede si no hay conexión o si el ticket ya existe y cómo "
          "se regresa desde el detalle. El flujo de la Figura 1 asegura la continuidad de la tarea y evita "
          "callejones sin salida."]),
        ("2. Explique la función de NavController y NavHost. ¿Cómo se representa el recorrido entre las pantallas de su proyecto?",
         ["NavController es el objeto que cambia de pantalla y guarda el historial en la pila de navegación: "
          "navigate() abre una pantalla encima de la pila y popBackStack() retira la actual para volver a la "
          "anterior. NavHost es el contenedor que asocia cada ruta con la función que dibuja su pantalla y muestra "
          "la que corresponde, comenzando por un destino inicial.",
          "En la compra, NavHost parte de la ruta login y continúa por home, tickets y tickets/{ticketId}, todas "
          "declaradas en la sealed class Route. El recorrido se representa en la Figura 1: el ingreso navega a home "
          "retirando el login de la pila, la pestaña «Mis tickets» navega a tickets sin duplicar destinos y el "
          "detalle recibe el identificador del ticket en la ruta y regresa con popBackStack()."]),
        ("3. Seleccione dos eventos del usuario implementados y describa qué ocurre desde que se produce el evento hasta que la interfaz responde.",
         ["En el evento onValueChange del campo de código, Compose entrega el texto nuevo; sanitizeCode() descarta "
          "lo que no es número y limita a ocho dígitos, y el valor se guarda en el estado de la pantalla. Luego "
          "validateCode() calcula el mensaje de error y Compose vuelve a dibujar el campo, el contador o el mensaje "
          "en rojo, y decide si el botón «Ingresar» queda habilitado.",
          "En el evento onClick de «Confirmar pago», la aplicación procesa el pago de demostración con la selección "
          "revisada y reemplaza el contenido del diálogo por el resultado. Si el pago es exitoso, el ticket queda "
          "registrado y aparece en «Mis tickets»; si la emisión queda pendiente, el diálogo indica que no se debe "
          "pagar otra vez."]),
        ("4. ¿Qué tipo de feedback utiliza su aplicación? ¿Por qué resulta comprensible para el usuario?",
         ["SisCome emplea feedback inmediato junto a los campos mediante mensajes, contadores y ayudas; feedback de "
          "estado con opciones resaltadas, total actualizado y botones habilitados o deshabilitados; y avisos de "
          "error, pendiente y éxito en el diálogo de pago, como resume la Tabla 2.",
          "Resulta comprensible porque aparece junto al elemento que lo origina y en el momento oportuno, ya que los "
          "errores se muestran después de interactuar con el campo. Además usa lenguaje cotidiano que indica la "
          "causa y la acción a seguir, como «El código debe tener 8 dígitos (llevas 3)», y combina color con texto "
          "para no depender solo del color (Durán-Coronado et al., 2019)."]),
        ("5. Identifique un error que su aplicación previene y explique cómo lo hace.",
         ["La aplicación previene la compra duplicada de un mismo ticket. Cuando el estudiante ya posee el ticket de "
          "una fecha y servicio, la opción no ofrece el pago: muestra «Ya tienes este ticket» y un atajo hacia el "
          "próximo menú disponible. Ante un pago con emisión pendiente se conserva la referencia y se indica que no "
          "se debe pagar otra vez.",
          "También previene el envío de credenciales incompletas: el botón «Ingresar» permanece deshabilitado hasta "
          "que validateCode() y validatePassword() no devuelven errores, y sanitizeCode() impide escribir "
          "caracteres no numéricos. Así el error se evita antes de que ocurra."]),
        ("6. Describa un patrón de interacción aplicado en su proyecto y justifique por qué es adecuado para la tarea.",
         ["Se aplicó el patrón de formulario con confirmación: tras elegir el menú, el usuario abre el diálogo "
          "«Revisar y pagar», que resume fechas, servicios y total, y solo se cobra al pulsar «Confirmar pago». "
          "Después, el mismo diálogo comunica el resultado.",
          "Es adecuado porque el pago es una acción con consecuencias económicas y difícil de revertir. La "
          "confirmación otorga control al usuario y le permite detectar una selección equivocada antes de pagar. "
          "Se complementa con el patrón lista-detalle de «Mis tickets», que permite consultar el ticket comprado "
          "sin perder la lista."]),
    ]
    for q, answers in qa:
        heading(cell, q)
        for a in answers:
            text(cell, a)


def conclusiones(cell):
    clear_trailing(cell)
    text(cell, "Del desarrollo de la tarea «Comprar ticket» se desprenden las siguientes conclusiones:")
    items = [
        "Navigation Compose, con NavController, NavHost y las rutas cerradas de la sealed class Route, permitió "
        "conectar las cuatro rutas de la compra con avance y regreso predecibles y sin callejones sin salida.",
        "El manejo de la pila de navegación es una decisión de interacción y no solo técnica: retirar el login con "
        "popUpTo e inclusive evita volver al formulario tras autenticarse.",
        "Los eventos onValueChange, onCheckedChange y onClick, unidos al estado de Compose, producen respuestas "
        "inmediatas que mantienen informado al usuario en cada acción.",
        "La prevención de errores mediante validación en línea, filtro de entrada, botones deshabilitados y bloqueo "
        "de compra duplicada resultó más efectiva que informar el error después de que ocurre.",
        "El formulario con confirmación y la lista-detalle se ajustan a la compra de tickets porque otorgan control "
        "al usuario en una operación sensible como el pago.",
        "Separar la validación en funciones puras de LoginValidation.kt permitió verificarla con pruebas unitarias, "
        "y la matriz de pruebas confirmó el comportamiento esperado en todos los casos evaluados.",
    ]
    for i, t in enumerate(items, 1):
        text(cell, f"{i}. {t}")


def main():
    draw_diagram(DIAGRAM)
    doc = Document(str(TEMPLATE))
    missing = []
    ejercicios(find_cell(doc, "SOLUCIÓN DE EJERCICIOS"), missing)
    cuestionario(find_cell(doc, "SOLUCIÓN DEL CUESTIONARIO"))
    conclusiones(find_cell(doc, "CONCLUSIONES"))
    doc.save(str(OUTPUT))
    print(f"Generado: {OUTPUT}")
    print("Capturas faltantes:", ", ".join(missing) if missing else "ninguna")


if __name__ == "__main__":
    main()
